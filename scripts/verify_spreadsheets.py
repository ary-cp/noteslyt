"""Offline spreadsheet regressions; synthetic workbooks, no DB or provider calls."""
from pathlib import Path
import sys
import tempfile
import zipfile
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from app.ingestion.parsers_docs import parse_csv, parse_tsv_file, parse_xlsx


def check(label, condition):
    assert condition, label
    print(f"PASS: {label}")


def main():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Costs"
        ws.append(["Item", "Value"])
        ws.append(["Uncached", "=SUM(1,2)"])
        ws.append(["Zero", "=1-1"])
        ws.append(["Text", "a|b\n<literal> & text"])
        ws.append(["Boolean", False])
        wb.create_sheet("Empty")
        wb.save(path)
        # Excel/LibreOffice write cached values. Inject a legitimate cached zero
        # to check that it isn't mistaken for a missing result.
        with zipfile.ZipFile(path) as source:
            entries = [(i, source.read(i.filename)) for i in source.infolist()]
        with zipfile.ZipFile(path, "w") as target:
            for info, data in entries:
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>1-1</f><v></v>", b"<f>1-1</f><v>0</v>")
                target.writestr(info, data)
        parsed = parse_xlsx(path, None)
        text = parsed.units[0].text
        check("uncached formula is explicit, not silently blank", "=SUM(1,2)" in text and "not calculated" in text)
        check("warning identifies formula cell", len(parsed.warnings) == 1 and "B2" in parsed.warnings[0])
        check("cached zero preserved without formula fallback", "| Zero | 0 |" in text and "=1-1" not in text)
        check("pipe, newline and HTML-like cell text preserved", "a\\|b<br>&lt;literal&gt; &amp; text" in text)
        check("boolean false preserved", "| Boolean | False |" in text)
        check("empty sheets skipped", len(parsed.units) == 1)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Gaps"
        ws.cell(3, 1, "Name")
        for i in range(201):
            ws.cell(5 + 2 * i, 1, f"Record {i}")
        wb.save(path)
        parsed = parse_xlsx(path, None)
        check("row citations respect leading/internal blank rows", parsed.units[0].locator_value == "Gaps!5:103")
        check("last row citation matches worksheet", parsed.units[-1].locator_value == "Gaps!405:405")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Column"])
        ws.append(["Included"])
        ws.cell(5001, 1, "Beyond limit")
        ws.cell(2, 30).number_format = "0.00"
        wb.save(path)
        original = ReadOnlyWorksheet.iter_rows
        def bounded(self, *args, **kwargs):
            assert kwargs.get("max_row") == 5000, "iteration must be bounded before allocation"
            return original(self, *args, **kwargs)
        with patch.object(ReadOnlyWorksheet, "iter_rows", bounded):
            parsed = parse_xlsx(path, None)
        check("row limit applied before materializing workbook", "Beyond limit" not in parsed.units[0].text and any("5000" in w for w in parsed.warnings))
        check("format-only trailing columns omitted", parsed.units[0].text.splitlines()[0] == "| Column |")
        # Every open read-only workbook should be closed, including on Windows.
        path.unlink()
        csv_path = Path(tmp) / "test.csv"
        csv_path.write_text('\nname,notes\nAlpha,"first line\nsecond line"\n\nBeta,ok,extra\n', encoding="utf-8")
        parsed = parse_csv(csv_path, None)
        check("CSV quoted newlines retained", "first line\nsecond line" in parsed.units[0].text)
        check("CSV extra columns retained with labels", "Column 3: extra" in parsed.units[0].text)
        check("CSV locators count records including blank records", parsed.units[0].locator_value == "3:5")
        csv_path.write_text('\n\nname,notes\n', encoding="utf-8")
        check("CSV header-only locator preserves original row", parse_csv(csv_path, None).units[0].locator_value == "3:3")
        csv_path.write_text('name\tnotes\nAlpha\t"first line\nsecond line"\n', encoding="utf-8")
        parsed = parse_tsv_file(csv_path, None)
        check("TSV shares multiline fix", parsed.parser == "tsv" and "first line\nsecond line" in parsed.units[0].text)
    print("SPREADSHEETS VERIFIED")


if __name__ == "__main__":
    main()
